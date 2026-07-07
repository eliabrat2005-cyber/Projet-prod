"""
pages/repartition_transport.py
==============================
Repartition des couts de transport entre Ferment Station (Symbiose + Niko) et
Spraga (kombucha), au prorata du POIDS des produits de chaque commande.

Source du cout : la reconciliation transport (Pennylane -> commande EasyBeer).
Source des lignes/poids : le detail commande EasyBeer (barcode + poidsUnitaire).
Classification FS/Spraga : regle mots-cles + table override data/entites_produits.csv.

Cycle de vie (PRD sect. 6) :
  - PROVISOIRE : calcul auto, recalcule a chaque « Mettre a jour ».
  - VALIDEE    : fige par l'operateur (telle quelle ou apres correction) - jamais
                 ecrase par la synchro.
  - ANOMALIE   : poids nul / reference non classee / cout absent - a traiter.

La page LIT la base ; le seul appel API est « Mettre a jour » (job de synchro).
"""
from __future__ import annotations

import asyncio
import io
import logging
from datetime import date, timedelta

from nicegui import ui

from common.services import allocation_store
from common.services.allocation_sync import synchroniser
from core.allocation.classification import reload_overrides
from core.allocation.models import (
    ENTITY_FS,
    ENTITY_SPRAGA,
    STATUS_ANOMALIE,
    STATUS_PROVISOIRE,
    STATUS_VALIDEE,
)
from pages.auth import require_auth
from pages.theme import (
    COLORS,
    error_banner,
    kpi_card,
    page_layout,
    section_title,
)

_log = logging.getLogger("ferment.repartition_transport")

_ENTITY_LABEL = {ENTITY_FS: "Ferment Station", ENTITY_SPRAGA: "Spraga"}

# Badge statut d'allocation -> couleur Quasar.
_STATUT_BADGE_JS = r"""
    <q-td :props="props">
      <q-badge :color="{
        'VALIDEE':'green-6',
        'PROVISOIRE':'blue-grey-5',
        'ANOMALIE':'red-6'
      }[props.value] || 'grey-6'"
      :label="{'VALIDEE':'Validee (figee)','PROVISOIRE':'Provisoire','ANOMALIE':'Anomalie'}[props.value] || props.value" />
    </q-td>
"""


# ─── Format (virgule francaise) ──────────────────────────────────────────────
def _eur(v) -> str:
    return f"{v:,.2f} €".replace(",", " ").replace(".", ",") if v is not None else "—"


def _kg(v) -> str:
    return f"{v:,.1f} kg".replace(",", " ").replace(".", ",") if v is not None else "—"


def _pct(v) -> str:
    return f"{v * 100:.1f} %".replace(".", ",") if v is not None else "—"


def _f(v):
    """float depuis un Decimal/None (les NUMERIC PG arrivent en Decimal)."""
    return float(v) if v is not None else None


def _default_sync_range() -> tuple[str, str]:
    """12 derniers mois au format 'YYYY-MM'."""
    d1 = (date.today() - timedelta(days=365)).replace(day=1)
    return d1.strftime("%Y-%m"), date.today().strftime("%Y-%m")


# ─── Page ────────────────────────────────────────────────────────────────────
@ui.page("/repartition-transport")
def page_repartition_transport():
    user = require_auth()
    if not user:
        return
    tenant_id = user.get("tenant_id")
    email = user.get("email") or "?"

    with page_layout("Répartition transport", "call_split", "/repartition-transport"):
        ui.label(
            "Répartition du coût de transport de chaque commande entre Ferment "
            "Station (Symbiose + Niko) et Spraga (kombucha), au prorata du poids. "
            "Le coût vient de la réconciliation transport ; validez pour figer."
        ).classes("text-body2").style(f"color: {COLORS['ink2']}")

        # État local partagé entre les closures.
        state: dict = {"rows": []}

        # ── Barre d'actions ────────────────────────────────────────────────
        with ui.card().classes("w-full").props("flat bordered"):
            with ui.card_section().classes("q-pa-md column gap-3"):
                section_title("Paramètres", "tune")
                with ui.row().classes("w-full gap-3 items-end wrap"):
                    month_from = ui.select(options={}, label="mois de début").props(
                        "outlined dense"
                    ).classes("w-48")
                    month_to = ui.select(options={}, label="mois de fin").props(
                        "outlined dense"
                    ).classes("w-48")
                    status_filter = ui.select(
                        options={
                            "": "Tous les statuts",
                            STATUS_PROVISOIRE: "Provisoire",
                            STATUS_VALIDEE: "Validée (figée)",
                            STATUS_ANOMALIE: "Anomalie",
                        },
                        value="",
                        label="statut",
                    ).props("outlined dense").classes("w-48")
                    run_btn = ui.button("Mettre à jour", icon="sync").props(
                        "outline dense color=green-8"
                    )
                    with run_btn:
                        ui.tooltip(
                            "Recalcule les répartitions provisoires depuis la "
                            "réconciliation + le détail EasyBeer (12 derniers mois)"
                        )
                ui.label(
                    "« Mettre à jour » interroge EasyBeer (quelques minutes). Les "
                    "commandes déjà validées ne sont jamais recalculées. Ensuite, "
                    "changer de mois ou de statut est instantané (lecture base)."
                ).classes("text-caption").style(f"color: {COLORS['ink2']}")

        progress_box = ui.column().classes("w-full")
        kpis_box = ui.column().classes("w-full")
        table_box = ui.column().classes("w-full")

        # ── Synchro (seul appel API) ───────────────────────────────────────
        async def _sync():
            progress_box.clear()
            run_btn.disable()
            with progress_box, ui.card().classes("w-full").props("flat bordered"):
                with ui.card_section().classes("q-pa-md column gap-2"):
                    prog_label = ui.label("Synchronisation en cours…").classes(
                        "text-body2"
                    ).style(f"color: {COLORS['ink']}")
                    ui.linear_progress(value=0, show_value=False, size="10px").props(
                        "rounded color=green-8 indeterminate"
                    )
            pm_from, pm_to = _default_sync_range()

            def _work():
                return synchroniser(tenant_id, pm_from, pm_to)

            try:
                stats = await asyncio.to_thread(_work)
            except Exception as exc:  # noqa: BLE001
                _log.exception("Synchro répartition échouée")
                progress_box.clear()
                with progress_box:
                    error_banner(f"Échec de la synchronisation : {exc}")
                run_btn.enable()
                return
            progress_box.clear()
            run_btn.enable()
            ui.notify(
                f"Synchro terminée : {stats['ecrites']} écrites, "
                f"{stats['anomalies']} anomalie(s), "
                f"{stats['figees_ignorees']} figée(s) préservée(s).",
                type="positive",
            )
            _refresh_months(select_latest=True)

        run_btn.on_click(_sync)

        # ── Rendu KPIs + tableau ───────────────────────────────────────────
        def _render(rows: list[dict]):
            state["rows"] = rows
            kpis_box.clear()
            table_box.clear()

            total = sum((_f(r["transport_cost_total"]) or 0) for r in rows)
            cost_fs = sum((_f(r["cost_fs"]) or 0) for r in rows)
            cost_sp = sum((_f(r["cost_spraga"]) or 0) for r in rows)
            n_prov = sum(1 for r in rows if r["allocation_status"] == STATUS_PROVISOIRE)
            n_val = sum(1 for r in rows if r["allocation_status"] == STATUS_VALIDEE)
            n_ano = sum(1 for r in rows if r["allocation_status"] == STATUS_ANOMALIE)

            with kpis_box, ui.row().classes("w-full gap-3 wrap"):
                kpi_card("euro", "Coût transport total", _eur(total))
                kpi_card("science", "Part Ferment Station", _eur(cost_fs), COLORS["green"])
                kpi_card("sports_bar", "Part Spraga", _eur(cost_sp), COLORS["blue"])
                kpi_card("hourglass_empty", "Provisoires", str(n_prov), COLORS["ink2"])
                kpi_card("verified", "Validées (figées)", str(n_val), COLORS["green"])
                kpi_card("error_outline", "Anomalies", str(n_ano), COLORS["error"])

            if not rows:
                with table_box:
                    ui.label(
                        "Aucune répartition pour cette sélection. Cliquez "
                        "« Mettre à jour » (après avoir synchronisé la réconciliation)."
                    ).classes("text-body2 q-mt-md").style(f"color: {COLORS['ink2']}")
                return

            columns = [
                {"name": "order_number", "label": "N° cmd", "field": "order_number",
                 "align": "left", "sortable": True},
                {"name": "client", "label": "Client", "field": "client",
                 "align": "left", "sortable": True},
                {"name": "order_status", "label": "Cmd", "field": "order_status",
                 "align": "left"},
                {"name": "wfs", "label": "Poids FS", "field": "wfs", "align": "right"},
                {"name": "wsp", "label": "Poids Spraga", "field": "wsp", "align": "right"},
                {"name": "pct", "label": "Répartition", "field": "pct", "align": "right"},
                {"name": "total", "label": "Coût total", "field": "total", "align": "right"},
                {"name": "cfs", "label": "Coût FS", "field": "cfs", "align": "right"},
                {"name": "csp", "label": "Coût Spraga", "field": "csp", "align": "right"},
                {"name": "status", "label": "Allocation", "field": "status",
                 "align": "left", "sortable": True},
            ]
            table_rows = [_row(r) for r in rows]

            with table_box:
                with ui.row().classes("w-full items-center justify-between q-mt-sm"):
                    section_title(f"Commandes ({len(rows)})", "table_rows")
                    with ui.row().classes("gap-2"):
                        val_btn = ui.button(
                            "Valider la sélection", icon="verified",
                        ).props("outline dense color=green-8")
                        ui.button(
                            "Exporter Excel", icon="download", on_click=lambda: _export(rows),
                        ).props("outline dense color=grey-8")

                table = ui.table(
                    columns=columns, rows=table_rows, row_key="order_number",
                    selection="multiple", pagination={"rowsPerPage": 50},
                ).classes("w-full").props(
                    'flat bordered dense :rows-per-page-options="[25,50,100]"'
                )
                table.add_slot("body-cell-status", _STATUT_BADGE_JS)

                # Clic sur une ligne -> detail commande.
                by_num = {r["order_number"]: r for r in rows}
                table.on(
                    "rowClick",
                    lambda e: _open_detail(by_num.get(e.args[1]["order_number"])),
                )

                async def _mass_validate():
                    nums = [
                        r["order_number"] for r in table.selected
                        if r.get("status") == STATUS_PROVISOIRE
                    ]
                    if not nums:
                        ui.notify(
                            "Sélectionnez des commandes provisoires à valider.",
                            type="warning",
                        )
                        return
                    n = await asyncio.to_thread(
                        allocation_store.mass_validate, tenant_id, nums, by=email
                    )
                    ui.notify(f"{n} commande(s) figée(s).", type="positive")
                    table.selected = []
                    _reload_table()

                val_btn.on_click(_mass_validate)

        def _row(r: dict) -> dict:
            pfs, psp = _f(r["pct_fs"]), _f(r["pct_spraga"])
            rep = f"{_pct(pfs)} / {_pct(psp)}" if pfs is not None else "—"
            return {
                "order_number": r["order_number"],
                "client": r["order_client"] or "—",
                "order_status": r["order_status"] or "—",
                "wfs": _kg(_f(r["weight_fs_kg"])),
                "wsp": _kg(_f(r["weight_spraga_kg"])),
                "pct": rep,
                "total": _eur(_f(r["transport_cost_total"])),
                "cfs": _eur(_f(r["cost_fs"])),
                "csp": _eur(_f(r["cost_spraga"])),
                "status": r["allocation_status"],
                "manual": bool(r.get("is_manual_override")),
            }

        # ── Détail commande (lignes + calcul + édition + audit) ────────────
        def _open_detail(r: dict | None):
            if not r:
                return
            with ui.dialog() as dlg, ui.card().classes("w-full").style("max-width: 900px"):
                num = r["order_number"]
                with ui.row().classes("w-full items-center justify-between"):
                    section_title(f"Commande {num} — {r['order_client'] or ''}", "receipt_long")
                    ui.button(icon="close", on_click=dlg.close).props("flat round dense")

                if r.get("anomaly_reason"):
                    error_banner(f"Anomalie : {r['anomaly_reason']}")

                # Lignes classées.
                lines = r.get("lines_json") or []
                if isinstance(lines, str):
                    import json
                    lines = json.loads(lines)
                lcols = [
                    {"name": "prod", "label": "Produit", "field": "prod", "align": "left"},
                    {"name": "ent", "label": "Entité", "field": "ent", "align": "left"},
                    {"name": "qte", "label": "Qté", "field": "qte", "align": "right"},
                    {"name": "pu", "label": "Poids u.", "field": "pu", "align": "right"},
                    {"name": "pl", "label": "Poids ligne", "field": "pl", "align": "right"},
                    {"name": "bc", "label": "Code-barres", "field": "bc", "align": "left"},
                ]
                lrows = [
                    {
                        "prod": ln.get("product_name") or "—",
                        "ent": _ENTITY_LABEL.get(ln.get("entity"), ln.get("entity") or "?"),
                        "qte": ln.get("quantity"),
                        "pu": _kg(ln.get("unit_weight_kg")),
                        "pl": _kg(ln.get("line_weight_kg")),
                        "bc": ln.get("barcode") or "—",
                    }
                    for ln in lines
                ]
                ui.table(columns=lcols, rows=lrows, row_key="prod").classes(
                    "w-full"
                ).props("flat bordered dense")

                # Calcul pas à pas.
                with ui.column().classes("gap-1 q-mt-sm"):
                    ui.label(
                        f"Poids FS {_kg(_f(r['weight_fs_kg']))} + Spraga "
                        f"{_kg(_f(r['weight_spraga_kg']))} = "
                        f"{_kg(_f(r['weight_total_kg']))}"
                    ).classes("text-caption").style(f"color: {COLORS['ink2']}")
                    ui.label(
                        f"Coût {_eur(_f(r['transport_cost_total']))} → FS "
                        f"{_eur(_f(r['cost_fs']))} · Spraga {_eur(_f(r['cost_spraga']))}"
                    ).classes("text-body2").style(f"color: {COLORS['ink']}; font-weight:600")

                # Édition manuelle + validation.
                total = _f(r["transport_cost_total"]) or 0.0
                is_frozen = r["allocation_status"] == STATUS_VALIDEE
                is_anomaly = r["allocation_status"] == STATUS_ANOMALIE
                ui.separator().classes("q-my-sm")

                # Bandeau « figée » + bouton « Modifier » (révèle l'édition).
                if is_frozen:
                    with ui.row().classes("items-center gap-2"):
                        ui.icon("lock", size="sm").style(f"color: {COLORS['success']}")
                        ui.label(
                            "Commande figée (validée) - non recalculée par la synchro."
                        ).classes("text-caption").style(f"color: {COLORS['success']}")
                        ui.button(
                            "Modifier", icon="edit",
                            on_click=lambda: setattr(edit_box, "visible", True),
                        ).props("flat dense color=grey-8")

                # Zone d'édition : visible d'emblée si non figée, sinon via « Modifier ».
                edit_box = ui.column().classes("w-full gap-2")
                edit_box.visible = not is_frozen
                with edit_box:
                    section_title("Corriger" if is_anomaly else "Corriger / valider", "edit")
                    with ui.row().classes("w-full gap-3 items-end wrap"):
                        cfs_in = ui.number(
                            "Coût FS (€)", value=_f(r["cost_fs"]) or 0.0, format="%.2f",
                        ).props("outlined dense").classes("w-40")
                        csp_in = ui.number(
                            "Coût Spraga (€)", value=_f(r["cost_spraga"]) or 0.0, format="%.2f",
                        ).props("outlined dense").classes("w-40")
                        reason_in = ui.input("Motif").props("outlined dense").classes("w-56")

                    async def _validate_as_is():
                        n = await asyncio.to_thread(
                            allocation_store.validate, tenant_id, num, by=email,
                            reason=reason_in.value or None,
                        )
                        if n:
                            ui.notify("Commande figée.", type="positive")
                            dlg.close()
                            _reload_table()
                        else:
                            ui.notify("Impossible de valider (déjà figée ou anomalie).",
                                      type="warning")

                    async def _save_manual():
                        cfs = round(float(cfs_in.value or 0), 2)
                        csp = round(float(csp_in.value or 0), 2)
                        if abs((cfs + csp) - round(total, 2)) > 0.01:
                            ui.notify(
                                f"FS + Spraga doit égaler {_eur(total)} "
                                f"(actuel {_eur(cfs + csp)}).",
                                type="negative",
                            )
                            return
                        n = await asyncio.to_thread(
                            allocation_store.apply_manual, tenant_id, num,
                            cost_fs=cfs, cost_spraga=csp, by=email,
                            reason=reason_in.value or None,
                        )
                        if n:
                            ui.notify(
                                "Modification enregistrée." if is_frozen
                                else "Correction enregistrée et commande figée.",
                                type="positive",
                            )
                            dlg.close()
                            _reload_table()
                        else:
                            ui.notify("Échec de l'enregistrement.", type="warning")

                    with ui.row().classes("gap-2"):
                        # « Valider tel quel » n'a de sens que sur une provisoire.
                        if not is_frozen and not is_anomaly:
                            ui.button("Valider tel quel", icon="verified",
                                      on_click=_validate_as_is).props("color=green-8")
                        ui.button(
                            "Enregistrer" if is_frozen else "Corriger + valider",
                            icon="save", on_click=_save_manual,
                        ).props("outline color=green-8")

                # Journal d'audit.
                audit = allocation_store.list_audit(tenant_id, num)
                if audit:
                    ui.separator().classes("q-my-sm")
                    section_title("Journal", "history")
                    acols = [
                        {"name": "when", "label": "Date", "field": "when", "align": "left"},
                        {"name": "who", "label": "Par", "field": "who", "align": "left"},
                        {"name": "action", "label": "Action", "field": "action", "align": "left"},
                        {"name": "detail", "label": "Détail", "field": "detail", "align": "left"},
                    ]
                    arows = [
                        {
                            "when": (a["changed_at"].strftime("%d/%m/%Y %H:%M")
                                     if a.get("changed_at") else "—"),
                            "who": a.get("changed_by") or "—",
                            "action": a.get("action") or "—",
                            "detail": " ".join(
                                x for x in [a.get("field"), a.get("value_after"),
                                            (f"({a['reason']})" if a.get("reason") else "")]
                                if x
                            ) or "—",
                        }
                        for a in audit
                    ]
                    ui.table(columns=acols, rows=arows, row_key="when").classes(
                        "w-full"
                    ).props("flat bordered dense")
            dlg.open()

        # ── Chargement / filtres ───────────────────────────────────────────
        def _current_range():
            a, b = month_from.value, month_to.value
            if a and b and a > b:
                a, b = b, a
            return a, b

        def _reload_table():
            a, b = _current_range()
            st = status_filter.value or None
            try:
                rows = allocation_store.list_allocations(tenant_id, a, b, st)
            except Exception:  # noqa: BLE001
                _log.exception("Lecture des répartitions échouée")
                rows = []
            _render(rows)

        def _refresh_months(select_latest: bool = False):
            reload_overrides()  # recharge la table override si le CSV a changé
            try:
                months = allocation_store.list_period_months(tenant_id)
            except Exception:  # noqa: BLE001
                _log.exception("Lecture des mois échouée")
                months = []
            opts = {m: m for m in months}
            month_from.set_options(opts)
            month_to.set_options(opts)
            if not months:
                _render([])
                return
            latest = months[0]
            guard["busy"] = True
            if select_latest or month_from.value not in opts:
                month_from.set_value(latest)
            if select_latest or month_to.value not in opts:
                month_to.set_value(latest)
            guard["busy"] = False
            _reload_table()

        guard = {"busy": False}

        def _on_filter_change(_e):
            if not guard["busy"]:
                _reload_table()

        month_from.on_value_change(_on_filter_change)
        month_to.on_value_change(_on_filter_change)
        status_filter.on_value_change(_on_filter_change)

        # ── Export Excel ────────────────────────────────────────────────────
        def _export(rows: list[dict]):
            data = _build_xlsx(rows)
            ui.download.content(
                data, "repartition_transport.xlsx",
                media_type=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
            )

        _refresh_months(select_latest=True)


def _build_xlsx(rows: list[dict]) -> bytes:
    """Export comptable : une ligne par commande, virgule française."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    def _num(v):
        return f"{float(v):.2f}".replace(".", ",") if v is not None else ""

    wb = Workbook()
    ws = wb.active
    ws.title = "Répartition"
    headers = [
        "N° commande", "Client", "Statut commande", "N° pièce", "Mois",
        "Poids FS (kg)", "Poids Spraga (kg)", "Poids total (kg)",
        "% FS", "% Spraga", "Coût total (€)", "Coût FS (€)", "Coût Spraga (€)",
        "Statut allocation", "Manuel", "Anomalie",
    ]
    ws.append(headers)
    for r in rows:
        pfs = float(r["pct_fs"]) if r.get("pct_fs") is not None else None
        psp = float(r["pct_spraga"]) if r.get("pct_spraga") is not None else None
        ws.append([
            r["order_number"], r.get("order_client"), r.get("order_status"),
            r.get("invoice_number"), r.get("period_month"),
            _num(r.get("weight_fs_kg")), _num(r.get("weight_spraga_kg")),
            _num(r.get("weight_total_kg")),
            (f"{pfs * 100:.1f}".replace(".", ",") if pfs is not None else ""),
            (f"{psp * 100:.1f}".replace(".", ",") if psp is not None else ""),
            _num(r.get("transport_cost_total")), _num(r.get("cost_fs")),
            _num(r.get("cost_spraga")), r.get("allocation_status"),
            "oui" if r.get("is_manual_override") else "", r.get("anomaly_reason") or "",
        ])
    head_fill = PatternFill("solid", fgColor="15803D")
    for ci in range(1, len(headers) + 1):
        c = ws.cell(1, ci)
        c.fill = head_fill
        c.font = Font(bold=True, color="FFFFFF")
        ws.column_dimensions[chr(64 + ci) if ci <= 26 else "AA"].width = 16
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
